#!/usr/bin/env python3
"""
Build the college-town home sales teaching dataset.

Inputs, both produced by extract_query.sql and neither one published:
    raw_extract.csv   one row per property: location and assessor-style attributes
    raw_events.csv    every property's listing and sale history, one row per event

Output: college_town_home_sales.{csv,sqlite,xlsx,duckdb,parquet}

Deterministic. The same two inputs always produce the same rows and the same listing_id
values. With the same library versions, the CSV, SQLite and Excel files are also
byte-identical from one build to the next (the Excel file's internal timestamps are pinned
for this). DuckDB and Parquet hold identical values on every build. Nothing depends on row
order in the source database, the current date or a random seed.

Usage, from the package root or from anywhere else:
    python build/build_dataset.py
    python build/build_dataset.py --input X.csv --events Y.csv --outdir some/dir
    python build/build_dataset.py --check-rules      # run the date-rule examples only

Requires pandas and openpyxl. Optional: duckdb for the .duckdb file, and pyarrow (or
fastparquet, or duckdb) for the .parquet file. A missing optional package skips only its
own format, with a message. make_parquet.py writes those two from the published CSV.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import re
import sqlite3
import sys
import zipfile
from pathlib import Path

import pandas as pd

BUILD_VERSION = "1.0"
TABLE_NAME = "home_sales"
STEM = "college_town_home_sales"
HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------------------------------
# Campus landmarks
# --------------------------------------------------------------------------------------
# Each row is assigned the nearest campus by straight-line distance from its own latitude
# and longitude, never from its city label. Names are each institution's official form.
# Every coordinate below was taken from the source named beside it (checked 2026-10-02).
# Where no landmark-level source was reachable, the university's own reference point is
# used; those sit within 0.1 mile of the historic core, which is the rounding precision of
# miles_to_campus.
CAMPUSES = {
    "University of Illinois Urbana-Champaign":         (40.1059, -88.2272),      # Foellinger Auditorium, Main Quad (Wikipedia)
    "University of Michigan":                          (42.2767, -83.7387),      # The Diag (latitude.to)
    "University of Wisconsin-Madison":                 (43.0753, -89.4041),      # Bascom Hill (latitude.to)
    "Purdue University":                               (40.42864, -86.91378),    # Engineering Fountain (Wikipedia)
    "University of Florida":                           (29.648834, -82.34329),   # Century Tower (Wikipedia)
    "University of Georgia":                           (33.9558, -83.3745),      # North Campus, university reference point (latitude.to)
    "The University of North Carolina at Chapel Hill": (35.912056, -79.051250),  # Old Well (Wikipedia)
    "University of Virginia":                          (38.0350, -78.5050),      # Central Grounds, university reference point (latitude.to)
    "The University of Iowa":                          (41.661333, -91.535639),  # Old Capitol (Wikipedia)
    "The Pennsylvania State University":               (40.7960, -77.8627),      # Old Main, university reference point (latitude.to)
    "University of California, Davis":                 (38.53972, -121.74917),   # Shields Library, on the Quad (Wikipedia)
}
EARTH_RADIUS_MILES = 3958.7613

# --------------------------------------------------------------------------------------
# Date rule
# --------------------------------------------------------------------------------------
# sale_date, contract_date and list_date all come from each property's event history
# (raw_events.csv), not from the source's date_posted column. date_posted is the date the
# home went on the market for most rows, not the sale date, and is not used anywhere.
WITHDRAWN_EVENTS = {"listingRemoved", "listedForRent"}
CONTRACT_EVENTS = {"pendingSale", "contingent"}
MAX_SILENCE_DAYS = 365    # a listing with no recorded activity for longer than this had lapsed
MAX_ESCROW_DAYS = 365     # a contract older than this at closing did not produce this sale
FELL_THROUGH_DAYS = 90    # a withdrawal more than this long before closing means the deal died

# Source column -> published column.
RENAMES = {
    "address_city": "city",
    "address_state": "state",
    "address_zipcode": "zipcode",
    "address_county": "county",
    "price": "sale_price",
    "living_area": "living_area_sqft",
    "lot_size": "lot_size_sqft",
    "reso_stories": "stories",
    "reso_basement_yn": "has_basement",
    "reso_has_garage": "has_garage",
    "reso_garage_parking_capacity": "garage_spaces",
    "reso_has_cooling": "has_cooling",
    "reso_has_heating": "has_heating",
    "reso_has_fireplace": "has_fireplace",
    "reso_architectural_style": "architectural_style",
    "reso_roof_type": "roof_type",
    "reso_tax_annual_amount": "tax_annual_amount",
    "reso_tax_assessed_value": "tax_assessed_value",
    "reso_high_school_district": "school_district",
}

TEXT_COLS = [
    "listing_id", "address_street", "city", "state", "zipcode", "county", "county_fips",
    "parcel_number", "university", "home_type", "architectural_style", "roof_type",
    "school_district",
]
DATE_COLS = ["list_date", "contract_date", "sale_date"]
INT_COLS = [
    "sale_price", "beds", "living_area_sqft", "lot_size_sqft", "year_built", "stories",
    "garage_spaces", "tax_assessed_value",
]
FLOAT_COLS = ["latitude", "longitude", "miles_to_campus", "baths", "tax_annual_amount"]
# 0/1 integers in every format: SQLite has no boolean type, and 0/1 behaves the same in
# pandas, SQL and JavaScript.
FLAG_COLS = ["has_basement", "has_garage", "has_cooling", "has_heating", "has_fireplace"]

FINAL_COLS = [
    "listing_id",
    "address_street", "city", "state", "zipcode", "county", "county_fips",
    "parcel_number", "latitude", "longitude",
    "university", "miles_to_campus",
    "list_date", "contract_date", "sale_date", "sale_price",
    "home_type", "beds", "baths", "living_area_sqft", "lot_size_sqft",
    "year_built", "stories", "has_basement", "has_garage", "garage_spaces",
    "has_cooling", "has_heating", "has_fireplace", "architectural_style", "roof_type",
    "tax_annual_amount", "tax_assessed_value", "school_district",
]

# Same indexes in SQLite and DuckDB.
INDEXES = [("city", "state"), ("university",), ("zipcode",), ("sale_date",),
           ("sale_price",), ("parcel_number",), ("address_street", "city", "state")]

COLUMN_NOTES = {
    "listing_id": "Stable 12-character id. See build_dataset.py for the exact recipe.",
    "address_street": "Street address. Not unique: condo and apartment units share one.",
    "city": "City as reported by the listing source.",
    "state": "Two-letter state code.",
    "zipcode": "5-digit ZIP. Text, so leading zeros survive.",
    "county": "County name.",
    "county_fips": "5-digit county FIPS code. Text.",
    "parcel_number": "County parcel number. The field that tells units in one building apart.",
    "latitude": "Decimal degrees, as reported.",
    "longitude": "Decimal degrees, as reported.",
    "university": "Nearest of the 11 campuses, by straight-line distance from this property's coordinates. Not derived from the city label.",
    "miles_to_campus": "Straight-line miles from this property to that campus's landmark, to one decimal. Not driving distance.",
    "list_date": "Date the home was listed in the listing that led to this sale. A failed contract does not reset it. Blank when no such listing is on record.",
    "contract_date": "Date the sale went under contract (first pending or contingent status of the deal that closed). Blank when no contract status is on record.",
    "sale_date": "Date the sale closed.",
    "sale_price": "Sale price in US dollars. Includes some implausible values, see the data description sheet.",
    "home_type": "Property type as reported. Nearly always SINGLE_FAMILY, including for condo units.",
    "beds": "Bedroom count.",
    "baths": "Bathroom count. Halves allowed.",
    "living_area_sqft": "Finished living area in square feet.",
    "lot_size_sqft": "Lot size in square feet.",
    "year_built": "Year built.",
    "stories": "Number of stories.",
    "has_basement": "1 = yes, 0 = no, blank = not reported.",
    "has_garage": "1 = yes, 0 = no, blank = not reported.",
    "garage_spaces": "Garage parking capacity.",
    "has_cooling": "1 = yes, 0 = no, blank = not reported.",
    "has_heating": "1 = yes, 0 = no, blank = not reported.",
    "has_fireplace": "1 = yes, 0 = no, blank = not reported.",
    "architectural_style": "Free-text style label. Not standardized.",
    "roof_type": "Free-text roof material. Not standardized.",
    "tax_annual_amount": "Most recent annual property tax amount, US dollars.",
    "tax_assessed_value": "Most recent assessed value, US dollars. Not the sale price.",
    "school_district": "High school district. Populated on under half of rows, and the values are not standardized.",
}


# --------------------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------------------

def norm_text(value) -> str:
    """Uppercase, trim, and collapse internal whitespace. Used for keys only."""
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NA:
        return ""
    return " ".join(str(value).upper().split())


def zip5(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NA:
        return ""
    return str(value).strip()[:5]


def make_listing_id(row) -> str:
    """
    Deterministic 12-hex-character id.

    Key fields, joined by '|': normalized street, city, state, 5-digit ZIP, sale date as
    ISO text (empty when missing), sale price as an integer string (empty when missing),
    parcel number verbatim (empty when missing).

    48 bits of hash. At about 36,000 rows the chance of any collision is about 2 in a
    million, and the build asserts uniqueness anyway.

    Changing ANY part of this recipe, including the normalization or the source of
    sale_date, regenerates every id in the dataset.
    """
    parts = [
        norm_text(row["address_street"]),
        norm_text(row["city"]),
        norm_text(row["state"]),
        zip5(row["zipcode"]),
        row["sale_date"] if isinstance(row["sale_date"], str) else "",
        "" if pd.isna(row["sale_price"]) else str(int(row["sale_price"])),
        "" if pd.isna(row["parcel_number"]) else str(row["parcel_number"]).strip(),
    ]
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()[:12]


def to_native(value):
    """Convert a pandas/numpy scalar to a plain Python value, or None if missing."""
    if value is None or value is pd.NA:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item"):          # numpy scalar
        return value.item()
    return value


def log(step: str, n: int, note: str = "") -> None:
    print(f"  {step:<46} {n:>7,} rows   {note}")


# --------------------------------------------------------------------------------------
# dates
# --------------------------------------------------------------------------------------

def derive_dates(events: list[tuple[dt.date, str]]):
    """
    Turn one property's chronological event history into its published dates.

    Returns (list_date, contract_date, sale_date, list_reason, contract_reason), where
    each date is a datetime.date or None and each reason says why a date is or is not set.

    sale_date      the most recent 'sold' event.
    contract_date  the start of the last unbroken run of pending/contingent events before
                   that sale. Dropped if it is more than MAX_ESCROW_DAYS before closing, or
                   if a withdrawal more than FELL_THROUGH_DAYS before closing shows that the
                   deal fell through.
    list_date      the start of the listing that led to the contract (or to the sale, when
                   there is no contract). A withdrawal or a switch to a rental listing ends a
                   listing. Coming back on the market after a failed contract does not: the
                   clock keeps running from the original listing. A silence longer than
                   MAX_SILENCE_DAYS means the listing had lapsed, and the clock restarts at
                   the next listing, if there is one.

    Only events after the previous sale belong to this sale.
    """
    sold = [d for d, e in events if e == "sold"]
    if not sold:
        return None, None, None, "no_sale", "no_sale"
    sale = max(sold)
    prev = max((d for d in sold if d < sale), default=None)
    cycle = [(d, e) for d, e in events if d <= sale and (prev is None or d > prev) and e != "sold"]

    contract, run = None, None
    for d, e in cycle:
        if e in CONTRACT_EVENTS:
            run = run or d
        else:
            contract, run = (run or contract), None
    contract = run or contract
    c_reason = "ok"
    if contract is None:
        c_reason = "no_contract"
    elif (sale - contract).days > MAX_ESCROW_DAYS:
        contract, c_reason = None, "stale_contract"
    elif any(e in WITHDRAWN_EVENTS and contract < d < sale and (sale - d).days > FELL_THROUGH_DAYS
             for d, e in cycle):
        contract, c_reason = None, "fell_through"

    end = contract or sale
    period = [(d, e) for d, e in cycle if d <= end]
    if not any(e == "listedForSale" for _, e in period):
        return None, contract, sale, "no_listing", c_reason
    withdrawals = [i for i, (d, e) in enumerate(period) if e in WITHDRAWN_EVENTS and d < end]
    current = period[withdrawals[-1] + 1:] if withdrawals else period
    lists = [d for d, e in current if e == "listedForSale"]
    if not lists:
        return None, contract, sale, "withdrawn", c_reason
    span = sorted({d for d, _ in current if d >= min(lists)} | {end})
    gaps = [b for a, b in zip(span, span[1:]) if (b - a).days > MAX_SILENCE_DAYS]
    if gaps:
        lists = [d for d in lists if d >= gaps[-1]]
        if not lists:
            return None, contract, sale, "lapsed", c_reason
    return min(lists), contract, sale, "ok", c_reason


# Executable examples of the rule. `--check-rules` runs them; the build runs them first.
_D = dt.date.fromisoformat
RULE_EXAMPLES = [
    ("ordinary sale",
     [("2024-03-01", "listedForSale"), ("2024-03-09", "contingent"), ("2024-03-12", "pendingSale"),
      ("2024-04-15", "sold")], ("2024-03-01", "2024-03-09", "2024-04-15")),
    ("back on market after a failed contract keeps the original listing date",
     [("2024-07-24", "listedForSale"), ("2024-07-27", "pendingSale"), ("2024-07-29", "listedForSale"),
      ("2024-07-29", "pendingSale"), ("2024-08-02", "sold")], ("2024-07-24", "2024-07-29", "2024-08-02")),
    ("only events after the previous sale count",
     [("2021-05-21", "listedForSale"), ("2021-06-05", "pendingSale"), ("2021-07-26", "sold"),
      ("2023-10-16", "listedForSale"), ("2023-11-30", "pendingSale"), ("2023-12-22", "sold")],
     ("2023-10-16", "2023-11-30", "2023-12-22")),
    ("relisted after a withdrawal starts a new listing",
     [("2023-01-10", "listedForSale"), ("2023-04-01", "listingRemoved"), ("2023-06-01", "listedForSale"),
      ("2023-06-20", "pendingSale"), ("2023-07-25", "sold")], ("2023-06-01", "2023-06-20", "2023-07-25")),
    ("withdrawn and later sold off-market: no listing date",
     [("2024-03-10", "listedForSale"), ("2024-06-10", "listingRemoved"), ("2025-12-17", "sold")],
     (None, None, "2025-12-17")),
    ("listing lapsed: two silent years before the contract",
     [("2023-06-06", "listedForSale"), ("2025-05-29", "priceChange"), ("2025-06-04", "contingent"),
      ("2025-07-01", "sold")], (None, "2025-06-04", "2025-07-01")),
    ("contract fell through: withdrawn four months before closing",
     [("2024-03-11", "pendingSale"), ("2024-04-01", "listingRemoved"), ("2024-07-25", "sold")],
     (None, None, "2024-07-25")),
    ("stale contract: the last contract is six years before closing",
     [("2017-04-27", "listedForSale"), ("2017-05-12", "pendingSale"), ("2023-04-24", "sold")],
     (None, None, "2023-04-24")),
    ("sold with no listing on record, contract kept",
     [("2026-03-01", "pendingSale"), ("2026-03-13", "sold")], (None, "2026-03-01", "2026-03-13")),
]


def check_rules() -> None:
    for label, evs, expect in RULE_EXAMPLES:
        got = derive_dates([(_D(d), e) for d, e in evs])[:3]
        got = tuple(x.isoformat() if x else None for x in got)
        if got != expect:
            raise SystemExit(f"date rule example failed: {label}\n  expected {expect}\n  got      {got}")
    print(f"  date rule: all {len(RULE_EXAMPLES)} examples pass")


def load_events(path: Path) -> pd.DataFrame:
    ev = pd.read_csv(path, dtype={"property_id": "int64", "event_date": str, "event": str})
    ev["event_date"] = pd.to_datetime(ev["event_date"], format="%Y-%m-%d").dt.date
    rows = []
    for pid, g in ev.groupby("property_id", sort=True):     # file order is chronological
        L, C, S, lr, cr = derive_dates(list(zip(g["event_date"], g["event"])))
        rows.append((pid, *(x.isoformat() if x else None for x in (L, C, S)), lr, cr))
    out = pd.DataFrame(rows, columns=["property_id", *DATE_COLS, "list_reason", "contract_reason"])
    for c in DATE_COLS:
        out[c] = out[c].astype(object)
    return out


# --------------------------------------------------------------------------------------
# build steps
# --------------------------------------------------------------------------------------

def load_raw(path: Path, normalize_city: bool = True) -> pd.DataFrame:
    """Read the extract as text, then coerce. Keeps leading zeros intact."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[""])
    df = df.drop(columns=["date_posted"])          # listing date, not sale date: never used
    df = df.rename(columns=RENAMES)

    # Text hygiene, applied before anything else reads these values: trim whitespace, and
    # treat placeholder strings as missing. The source has a handful (parcel numbers 'NA',
    # 'NONE', 'unknown' and '.', roof types 'None' and 'Unknown', 9 school districts 'NA').
    # Left as text they would reach the CSV, where pandas reads a literal NA back as missing
    # anyway, so the file would disagree with itself depending on the reader.
    placeholders = {"NA", "N/A", "NONE", "NULL", "UNKNOWN", "--", "-", "."}
    for col in [c for c in TEXT_COLS if c in df.columns]:
        df[col] = df[col].astype("string").str.strip()
        df[col] = df[col].mask(df[col].str.upper().isin(placeholders))
        df[col] = df[col].mask(df[col].eq(""))

    if normalize_city:
        # A few cities are spelled inconsistently ('URBANA' beside 'Urbana'), which silently
        # splits any GROUP BY city. Each spelling group collapses to its most common form;
        # ties go to the alphabetically first spelling so the result never depends on order.
        canonical = (
            df.assign(_k=df["city"].str.upper())
              .groupby("_k")["city"]
              .agg(lambda s: sorted(s.value_counts().items(), key=lambda kv: (-kv[1], kv[0]))[0][0])
        )
        changed = int((df["city"] != df["city"].str.upper().map(canonical)).sum())
        df["city"] = df["city"].str.upper().map(canonical)
        if changed:
            print(f"  (city spelling normalized on {changed} rows)")

    for col in INT_COLS:
        df[col] = pd.to_numeric(df[col], errors="raise").astype("Int64")
    for col in ["latitude", "longitude", "baths", "tax_annual_amount"]:
        df[col] = pd.to_numeric(df[col], errors="raise")
    for col in FLAG_COLS:
        mapped = df[col].map({"true": 1, "false": 0})
        if int((df[col].notna() & mapped.isna()).sum()):
            raise SystemExit(f"unexpected value in {col}: {sorted(set(df[col].dropna()) - {'true', 'false'})}")
        df[col] = mapped.astype("Int64")
    df["property_id"] = pd.to_numeric(df["property_id"], errors="raise").astype("int64")
    return df


def attach_dates(df: pd.DataFrame, dates: pd.DataFrame) -> pd.DataFrame:
    out = df.merge(dates, on="property_id", how="left", validate="one_to_one")
    for c in DATE_COLS:
        out[c] = out[c].astype(object).where(out[c].notna(), None)
    out["list_reason"] = out["list_reason"].fillna("no_sale")
    out["contract_reason"] = out["contract_reason"].fillna("no_sale")
    return out


def add_keys(df: pd.DataFrame) -> pd.DataFrame:
    df["_street_key"] = df["address_street"].map(norm_text)
    df["_city_key"] = df["city"].map(norm_text)
    df["_state_key"] = df["state"].map(norm_text)
    df["_zip_key"] = df["zipcode"].map(zip5)
    df["_parcel_key"] = df["parcel_number"].map(lambda v: "" if v is None or v is pd.NA else str(v).strip())
    return df


ADDR_KEY = ["_street_key", "_city_key", "_state_key", "_zip_key"]


def drop_duplicate_properties(df: pd.DataFrame) -> pd.DataFrame:
    """
    One row per property.

    Rows that share a street address AND a parcel number are one property listed twice in
    the source under two ids. Two cases:

    * Same transaction recorded twice (prices within 1% and sale dates within 90 days):
      keep the most complete record, then the later date, then the lower source id.
    * A genuine resale of the property: keep the most recent sale.

    Rows that share only a street address are separate units in one building (each has its
    own parcel number) and are all kept. In this source every repeated address carries a
    parcel number; a repeated address with a blank parcel would also be kept.
    """
    has_parcel = df["_parcel_key"] != ""
    key = ADDR_KEY + ["_parcel_key"]
    size = df[has_parcel].groupby(key)["address_street"].transform("size")
    dup = df.loc[size[size > 1].index]
    if dup.empty:
        log("2. one row per property", len(df), "(no duplicates)")
        return df

    attrs = [c for c in FINAL_COLS if c in df.columns]
    completeness = df[attrs].notna().sum(axis=1)
    drop, same_tx, resale = [], 0, 0
    for _, g in dup.groupby(key, sort=True):
        prices = g["sale_price"].dropna()
        dates = pd.to_datetime(g["sale_date"].dropna())
        same = (len(prices) == len(g) and len(dates) == len(g)
                and prices.max() <= prices.min() * 1.01
                and (dates.max() - dates.min()).days <= 90)
        g = g.assign(_c=completeness.loc[g.index], _d=g["sale_date"].fillna(""))
        if same:
            same_tx += 1
            order = g.sort_values(["_c", "_d", "property_id"], ascending=[False, False, True], kind="mergesort")
        else:
            resale += 1
            order = g.sort_values(["_d", "_c", "property_id"], ascending=[False, False, True], kind="mergesort")
        drop.extend(order.index[1:])
    out = df.drop(index=drop)
    log("2. one row per property", len(out),
        f"(-{len(drop)}: {same_tx} sales recorded twice, {resale} resold)")
    return out


def drop_small_cities(df: pd.DataFrame, min_rows: int) -> pd.DataFrame:
    """
    Keep only cities with at least `min_rows` sales.

    The collection reached into surrounding towns, leaving a long tail of places with a
    handful of sales each. Those rows carry most of the junk city labels and are too thin to
    analyse, so the tail is cut rather than explained away. --min-city-rows 0 keeps all.
    """
    if min_rows <= 0:
        return df
    counts = df.groupby(["_city_key", "_state_key"])["address_street"].transform("size")
    keep = counts >= min_rows
    places = df.loc[~keep].groupby(["_city_key", "_state_key"]).ngroups
    out = df[keep].copy()
    log(f"3. dropped cities under {min_rows} sales", len(out),
        f"(-{len(df) - len(out)} rows, {places} places)")
    return out


def add_campus(df: pd.DataFrame) -> pd.DataFrame:
    """Nearest campus by great-circle distance from the row's own coordinates."""
    import numpy as np

    if df["latitude"].isna().any() or df["longitude"].isna().any():
        raise SystemExit("rows without coordinates cannot be assigned a campus")
    lat, lon = np.radians(df["latitude"].to_numpy()), np.radians(df["longitude"].to_numpy())
    names = list(CAMPUSES)
    dists = np.column_stack([
        2 * EARTH_RADIUS_MILES * np.arcsin(np.sqrt(
            np.sin((np.radians(clat) - lat) / 2) ** 2
            + np.cos(lat) * np.cos(np.radians(clat)) * np.sin((np.radians(clon) - lon) / 2) ** 2
        ))
        for clat, clon in CAMPUSES.values()
    ])
    df["university"] = [names[i] for i in dists.argmin(axis=1)]
    df["miles_to_campus"] = dists.min(axis=1).round(1)
    log("4. assigned nearest campus", len(df),
        f"({df['university'].nunique()} campuses, max {df['miles_to_campus'].max():.1f} mi)")
    return df


def add_listing_id(df: pd.DataFrame) -> pd.DataFrame:
    df["listing_id"] = df.apply(make_listing_id, axis=1)
    dupes = int(df["listing_id"].duplicated().sum())
    if dupes:
        raise SystemExit(f"FATAL: {dupes} duplicate listing_id values. Widen the hash key.")
    log("5. generated listing_id", len(df), "(all unique)")
    return df


def normalize_types(df: pd.DataFrame) -> pd.DataFrame:
    """
    Published columns in a fixed order with fixed types, whatever the pandas version:
    nullable Int64 for integers and flags, float64 for decimals, and plain Python str (None
    when missing) for text and for the ISO dates. make_parquet.py calls this too, so a file
    written from the CSV matches one written by a full rebuild.
    """
    out = df[FINAL_COLS].copy()
    for c in INT_COLS + FLAG_COLS:
        out[c] = pd.to_numeric(out[c], errors="raise").astype("Int64")
    for c in FLOAT_COLS:
        out[c] = pd.to_numeric(out[c], errors="raise").astype("float64")
    for c in TEXT_COLS + DATE_COLS:
        s = out[c].astype(object)
        out[c] = s.where(s.notna(), None)
    return out.reset_index(drop=True)


def finalize(df: pd.DataFrame) -> pd.DataFrame:
    df = df.sort_values(["state", "city", "address_street", "sale_date", "listing_id"],
                        kind="mergesort", na_position="last")
    return normalize_types(df)


# --------------------------------------------------------------------------------------
# writers
# --------------------------------------------------------------------------------------

def sqlite_type(col: str) -> str:
    if col in INT_COLS or col in FLAG_COLS:
        return "INTEGER"
    if col in FLOAT_COLS:
        return "REAL"
    return "TEXT"           # text, and ISO dates: SQLite has no date type


def duckdb_type(col: str) -> str:
    if col in INT_COLS or col in FLAG_COLS:
        return "BIGINT"
    if col in FLOAT_COLS:
        return "DOUBLE"
    if col in DATE_COLS:
        return "DATE"
    return "VARCHAR"


def write_csv(df: pd.DataFrame, outdir: Path) -> Path:
    path = outdir / f"{STEM}.csv"
    df.to_csv(path, index=False)
    return path


def write_sqlite(df: pd.DataFrame, outdir: Path) -> Path:
    path = outdir / f"{STEM}.sqlite"
    if path.exists():
        path.unlink()
    cols_ddl = [f"    {c} {sqlite_type(c)}" + (" PRIMARY KEY" if c == "listing_id" else "")
                for c in FINAL_COLS]
    con = sqlite3.connect(path)
    try:
        cur = con.cursor()
        cur.execute(f"CREATE TABLE {TABLE_NAME} (\n" + ",\n".join(cols_ddl) + "\n)")
        # to_native is essential: a numpy int64 supports the buffer protocol, so sqlite3
        # would store it as a BLOB rather than an INTEGER, and every aggregate over that
        # column would silently return nonsense.
        records = [tuple(to_native(v) for v in rec) for rec in df.itertuples(index=False, name=None)]
        cur.executemany(f"INSERT INTO {TABLE_NAME} VALUES ({','.join('?' * len(FINAL_COLS))})", records)
        for cols in INDEXES:
            cur.execute(f"CREATE INDEX idx_{'_'.join(cols)} ON {TABLE_NAME} ({', '.join(cols)})")
        con.commit()
        cur.execute("VACUUM")
        con.commit()
    finally:
        con.close()
    return path


def as_dates(df: pd.DataFrame) -> pd.DataFrame:
    """
    Copy of df with the date columns as datetime.date objects and None for missing.

    Everywhere else in this build the dates are ISO text. Handing that text to Parquet or
    DuckDB would store a string column; converting here keeps both on a real date type.
    The .where() matters: .dt.date on its own leaves pandas NaT in an object column, and
    Arrow will not turn NaT into a date32 null.
    """
    out = df.copy()
    for c in DATE_COLS:
        parsed = pd.to_datetime(out[c], format="%Y-%m-%d", errors="raise")
        out[c] = parsed.dt.date.astype(object).where(parsed.notna(), None)
    return out


def write_duckdb(df: pd.DataFrame, outdir: Path) -> Path | None:
    """Same table, primary key and indexes as the SQLite file, with real DATE columns."""
    try:
        import duckdb
    except ImportError:
        print("  DuckDB skipped: install duckdb (pip install duckdb), then re-run.")
        return None
    path = outdir / f"{STEM}.duckdb"
    if path.exists():
        path.unlink()
    cols_ddl = ",\n".join(f"    {c} {duckdb_type(c)}" for c in FINAL_COLS)
    con = duckdb.connect(str(path))
    try:
        con.execute(f"CREATE TABLE {TABLE_NAME} (\n{cols_ddl},\n    PRIMARY KEY (listing_id)\n)")
        con.register("staged", as_dates(df))
        con.execute(f"INSERT INTO {TABLE_NAME} SELECT {', '.join(FINAL_COLS)} FROM staged")
        con.unregister("staged")
        for cols in INDEXES:
            con.execute(f"CREATE INDEX idx_{'_'.join(cols)} ON {TABLE_NAME} ({', '.join(cols)})")
    finally:
        con.close()
    return path


def write_parquet(df: pd.DataFrame, outdir: Path) -> Path | None:
    path = outdir / f"{STEM}.parquet"
    staged = as_dates(df)
    for engine in ("pyarrow", "fastparquet"):
        try:
            __import__(engine)
        except ImportError:
            continue
        staged.to_parquet(path, engine=engine, index=False, compression="zstd")
        return path
    try:
        import duckdb       # DuckDB can write Parquet without pyarrow installed
    except ImportError:
        print("  Parquet skipped: install pyarrow or duckdb (pip install pyarrow), then re-run.")
        return None
    con = duckdb.connect()
    try:
        con.register("staged", staged)
        con.execute(f"COPY (SELECT {', '.join(FINAL_COLS)} FROM staged) "
                    f"TO '{path.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)")
    finally:
        con.close()
    return path


PINNED_TIMESTAMP = "2026-03-26T00:00:00Z"   # end of collection
PINNED_ZIP_TIME = (2026, 3, 26, 0, 0, 0)


def _pin_xlsx_timestamps(path: Path) -> None:
    """
    openpyxl stamps the save time into docProps/core.xml and into every zip entry, so two
    builds of identical data would differ. Rewrite both with a fixed time.
    """
    with zipfile.ZipFile(path) as src:
        parts = [(i.filename, src.read(i.filename)) for i in src.infolist()]
    stamp = PINNED_TIMESTAMP.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)",
                              rb"\g<1>" + stamp + rb"\g<2>", data)
            info = zipfile.ZipInfo(name, date_time=PINNED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            out.writestr(info, data)


def write_xlsx(df: pd.DataFrame, outdir: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = outdir / f"{STEM}.xlsx"
    wb = Workbook()
    wb.properties.creator = "college-town home sales build"
    ws = wb.active
    ws.title = "home_sales"
    ws.append(FINAL_COLS)

    xl = df.copy()
    for c in DATE_COLS:                     # real dates in Excel, ISO text everywhere else
        xl[c] = pd.to_datetime(xl[c], format="%Y-%m-%d", errors="raise")
    for rec in xl.itertuples(index=False, name=None):
        ws.append([v.to_pydatetime().date() if isinstance(v, pd.Timestamp) else to_native(v)
                   for v in rec])

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {"listing_id": 14, "address_street": 28, "city": 16, "county": 20, "parcel_number": 18,
              "university": 42, "miles_to_campus": 15, "list_date": 12, "contract_date": 13,
              "sale_date": 12, "sale_price": 12, "living_area_sqft": 16, "lot_size_sqft": 14,
              "architectural_style": 20, "roof_type": 16, "tax_annual_amount": 16,
              "tax_assessed_value": 17, "school_district": 18}
    for i, col in enumerate(FINAL_COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    for c in DATE_COLS:
        for cell in ws[get_column_letter(FINAL_COLS.index(c) + 1)][1:]:
            cell.number_format = "yyyy-mm-dd"

    notes = wb.create_sheet("columns")
    notes.append(["column", "type", "non-empty", "notes"])
    for cell in notes[1]:
        cell.font, cell.fill = header_font, header_fill
    for col in FINAL_COLS:
        filled = int(df[col].notna().sum())
        kind = ("integer" if col in INT_COLS else "number" if col in FLOAT_COLS else
                "flag 0/1" if col in FLAG_COLS else "date" if col in DATE_COLS else "text")
        notes.append([col, kind, f"{filled:,} ({filled / len(df):.1%})", COLUMN_NOTES[col]])
    notes.freeze_panes = "A2"
    for letter, width in zip("ABCD", (22, 10, 18, 100)):
        notes.column_dimensions[letter].width = width
    for row in notes.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
    _pin_xlsx_timestamps(path)
    return path


# --------------------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=HERE / "raw_extract.csv", type=Path)
    ap.add_argument("--events", default=HERE / "raw_events.csv", type=Path)
    ap.add_argument("--outdir", default=HERE.parent / "data", type=Path)
    ap.add_argument("--keep-city-variants", action="store_true",
                    help="do not normalize inconsistent city spellings")
    ap.add_argument("--min-city-rows", type=int, default=90,
                    help="drop cities with fewer sales than this (0 keeps every city)")
    ap.add_argument("--check-rules", action="store_true", help="run the date-rule examples and exit")
    args = ap.parse_args()

    print(f"Building college-town home sales dataset (build {BUILD_VERSION})")
    check_rules()
    if args.check_rules:
        return 0
    for p in (args.input, args.events):
        if not p.exists():
            raise SystemExit(f"input not found: {p}")
    args.outdir.mkdir(parents=True, exist_ok=True)

    dates = load_events(args.events)
    df = load_raw(args.input, normalize_city=not args.keep_city_variants)
    df = attach_dates(df, dates)
    log("1. loaded extract and event history", len(df),
        f"(sale date on {df['sale_date'].notna().sum():,})")
    df = add_keys(df)
    df = drop_duplicate_properties(df)
    df = drop_small_cities(df, args.min_city_rows)
    df = add_campus(df)
    df = add_listing_id(df)
    print("  date coverage after filtering:")
    for c, r in (("list_date", "list_reason"), ("contract_date", "contract_reason")):
        print(f"    {c:<14} {df[c].notna().sum():>7,} set   blank: "
              + ", ".join(f"{k} {v:,}" for k, v in df.loc[df[c].isna(), r].value_counts().items()))
    out = finalize(df)
    log("6. final dataset", len(out), f"({len(FINAL_COLS)} columns)")

    print("\nWriting files:")
    written = [write_csv(out, args.outdir), write_sqlite(out, args.outdir), write_xlsx(out, args.outdir)]
    written += [p for p in (write_duckdb(out, args.outdir), write_parquet(out, args.outdir)) if p]
    for path in written:
        print(f"  {path.name:<40} {path.stat().st_size / 1_048_576:>7.2f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
